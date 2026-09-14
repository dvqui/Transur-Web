"""
Servicio de exportación a Excel (.xlsx) con formato profesional, generado
EN MEMORIA (BytesIO) para servirse como descarga directa desde Streamlit
con st.download_button (no se escribe nada en el disco del servidor).
"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import COLORS

AZUL_HEX = COLORS["azul_profundo"].replace("#", "")


def generar_reporte_excel(
    titulo: str, encabezados: list, filas: list, columnas_moneda: list = None,
    nombre_hoja: str = "Reporte",
) -> bytes:
    columnas_moneda = columnas_moneda or []

    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja
    num_cols = len(encabezados)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    celda_titulo = ws.cell(row=1, column=1, value=titulo)
    celda_titulo.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    celda_titulo.alignment = Alignment(horizontal="center", vertical="center")
    celda_titulo.fill = PatternFill(start_color=AZUL_HEX, end_color=AZUL_HEX, fill_type="solid")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
    fecha_str = (
        f"Cooperativa de Transportes 'Transur 7 de Mayo'  |  Generado: "
        f"{datetime.now().strftime('%d/%m/%Y %H:%M')}"
    )
    celda_sub = ws.cell(row=2, column=1, value=fecha_str)
    celda_sub.font = Font(name="Calibri", size=10, italic=True, color="5A6B7B")
    celda_sub.alignment = Alignment(horizontal="center")

    fila_encabezado = 4
    borde = Border(
        left=Side(style="thin", color="D0D0D0"), right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"), bottom=Side(style="thin", color="D0D0D0"),
    )

    for col_idx, nombre_col in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_encabezado, column=col_idx, value=nombre_col)
        celda.font = Font(bold=True, color="FFFFFF", size=11)
        celda.fill = PatternFill(start_color=AZUL_HEX, end_color=AZUL_HEX, fill_type="solid")
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde

    for fila_idx, fila_datos in enumerate(filas, start=fila_encabezado + 1):
        color_fondo = "F3F6F9" if fila_idx % 2 == 0 else "FFFFFF"
        for col_idx, valor in enumerate(fila_datos, start=1):
            celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
            celda.border = borde
            celda.fill = PatternFill(start_color=color_fondo, end_color=color_fondo, fill_type="solid")
            if (col_idx - 1) in columnas_moneda and isinstance(valor, (int, float)):
                celda.number_format = '"$"#,##0.00'
                celda.alignment = Alignment(horizontal="right")
            else:
                celda.alignment = Alignment(horizontal="left", vertical="center")

    for col_idx in range(1, num_cols + 1):
        letra = get_column_letter(col_idx)
        max_len = len(str(encabezados[col_idx - 1]))
        for fila_datos in filas:
            valor = fila_datos[col_idx - 1] if col_idx - 1 < len(fila_datos) else ""
            max_len = max(max_len, len(str(valor)))
        ws.column_dimensions[letra].width = max_len + 4

    ws.freeze_panes = f"A{fila_encabezado + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
